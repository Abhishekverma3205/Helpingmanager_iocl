"""
Fuel Station Manager — Flask backend
A single-station operations system: customer credit (udhari) management, staff &
attendance, daily sales/testing/tank-dip logging, and owner/manager/staff
authentication. Currency, locale and business details are all configurable per
deployment from Settings, so the same codebase serves any country/market.
"""
import os
import io
import csv
import re
import math
import secrets
from datetime import datetime, timedelta
from functools import wraps

try:
    # Optional convenience for local/dev runs: load variables from a .env
    # file if python-dotenv is installed. Never overrides variables the
    # platform/host has already set in the real environment.
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

from flask import Flask, request, session, jsonify, render_template, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.middleware.proxy_fix import ProxyFix
from openpyxl import Workbook, load_workbook

import database as db

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
IS_PRODUCTION = os.environ.get("FLASK_ENV", "development").lower() == "production"
SECRET_KEY = os.environ.get("SECRET_KEY")
if not SECRET_KEY:
    if IS_PRODUCTION:
        raise RuntimeError(
            "SECRET_KEY environment variable is required when FLASK_ENV=production. "
            "Generate one with: python -c \"import secrets; print(secrets.token_hex(32))\""
        )
    SECRET_KEY = "dev-secret-change-me"

app = Flask(__name__)
app.config["SECRET_KEY"] = SECRET_KEY
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = IS_PRODUCTION  # only send the cookie over HTTPS in production
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(hours=12)
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024  # 8 MB upload cap
# Trust one hop of X-Forwarded-* headers when running behind a reverse proxy / PaaS router
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

FUEL_TYPES = ["diesel", "petrol", "cng"]
DIP_FUEL_TYPES = ["diesel", "petrol"]  # only these two have physical tanks with dip charts
SHIFTS = ["morning", "evening", "night"]
ATTENDANCE_STATUSES = ["present", "absent", "half_day", "leave", "off"]
ROLE_RANK = {"staff": 1, "manager": 2, "admin": 3}
LOGIN_MAX_ATTEMPTS = 5
LOGIN_LOCKOUT_MINUTES = 15
CSRF_EXEMPT_ENDPOINTS = {
    "login", "index", "static", "register_owner", "register_manager",
    "pump_status", "csrf_token",
}
PUMP_CODE_RE = re.compile(r"^\d{6}$")


# ---------------------------------------------------------------------------
# Security: headers, CSRF, login rate limiting
# ---------------------------------------------------------------------------
@app.after_request
def set_security_headers(resp):
    resp.headers["X-Content-Type-Options"] = "nosniff"
    resp.headers["X-Frame-Options"] = "DENY"
    resp.headers["Referrer-Policy"] = "same-origin"
    resp.headers["Permissions-Policy"] = "geolocation=(), microphone=(), camera=()"
    if IS_PRODUCTION:
        resp.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    return resp


def get_csrf_token():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_hex(24)
    return session["csrf_token"]


@app.before_request
def enforce_csrf():
    """Every state-changing request from an authenticated session must carry
    the X-CSRF-Token header issued at login, so a malicious third-party site
    can't ride the user's session cookie to modify data."""
    if request.method not in ("POST", "PUT", "DELETE", "PATCH"):
        return
    if request.endpoint in CSRF_EXEMPT_ENDPOINTS:
        return
    if "user_id" not in session:
        return  # login_required/roles_required will reject these anyway
    sent = request.headers.get("X-CSRF-Token", "")
    expected = session.get("csrf_token", "")
    if not expected or not secrets.compare_digest(sent, expected):
        return jsonify({"error": "Your session has expired or is invalid. Please refresh and try again."}), 403


def _client_ip():
    return request.headers.get("X-Forwarded-For", request.remote_addr or "").split(",")[0].strip()


def is_login_locked(identifier):
    row = db.query_db("SELECT * FROM login_attempts WHERE identifier=?", (identifier,), one=True)
    if not row or not row["locked_until"]:
        return None
    locked_until = datetime.fromisoformat(row["locked_until"])
    if datetime.now() < locked_until:
        return locked_until
    return None


def register_failed_login(identifier):
    row = db.query_db("SELECT * FROM login_attempts WHERE identifier=?", (identifier,), one=True)
    count = (row["failed_count"] if row else 0) + 1
    locked_until = None
    if count >= LOGIN_MAX_ATTEMPTS:
        locked_until = (datetime.now() + timedelta(minutes=LOGIN_LOCKOUT_MINUTES)).isoformat()
        count = 0
    db.execute_db(
        "INSERT INTO login_attempts (identifier,failed_count,locked_until) VALUES (?,?,?) "
        "ON CONFLICT(identifier) DO UPDATE SET failed_count=excluded.failed_count, locked_until=excluded.locked_until",
        (identifier, count, locked_until),
    )


def clear_failed_logins(identifier):
    db.execute_db("DELETE FROM login_attempts WHERE identifier=?", (identifier,))


# ---------------------------------------------------------------------------
# Auth helpers
# ---------------------------------------------------------------------------
def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if "user_id" not in session:
            return jsonify({"error": "Authentication required"}), 401
        return f(*args, **kwargs)
    return wrapper


def roles_required(*roles):
    def decorator(f):
        @wraps(f)
        def wrapper(*args, **kwargs):
            if "user_id" not in session:
                return jsonify({"error": "Authentication required"}), 401
            if session.get("role") not in roles:
                return jsonify({"error": "You do not have permission to do this"}), 403
            return f(*args, **kwargs)
        return wrapper
    return decorator


def current_user():
    if "user_id" not in session:
        return None
    return db.query_db(
        "SELECT id,name,email,role,avatar,personal_id FROM users WHERE id=?", (session["user_id"],), one=True
    )


def serialize_user(u):
    if not u:
        return None
    out = dict(u)
    out["role_label"] = db.ROLE_LABELS.get(out.get("role"), out.get("role"))
    out.pop("password_hash", None)
    return out


def audit(action, entity, entity_id=None, details=None, success=True, actor_override=None):
    u = actor_override or current_user()
    db.log_action(
        u["id"] if u else None, u["name"] if u else None,
        action, entity, entity_id, details, _client_ip(), success=success,
    )


def recompute_status(amount, paid_amount):
    if paid_amount >= amount - 0.001:
        return "paid"
    if paid_amount > 0:
        return "partial"
    return "pending"


def get_setting(key, default=None):
    row = db.query_db("SELECT value FROM settings WHERE key=?", (key,), one=True)
    return row["value"] if row else default


def dip_volume_ltr(dip_mm, diameter_m, length_m):
    """Litres of liquid in a horizontal cylindrical tank given a dip
    (liquid height from the bottom, in mm), via the circular-segment formula."""
    r = diameter_m / 2.0
    h = max(0.0, min(dip_mm / 1000.0, diameter_m))  # clamp to [0, diameter]
    if h <= 0:
        return 0.0
    if h >= diameter_m:
        return round(math.pi * r * r * length_m * 1000, 2)
    segment_area = (r ** 2) * math.acos((r - h) / r) - (r - h) * math.sqrt(max(2 * r * h - h ** 2, 0))
    volume_m3 = segment_area * length_m
    return round(volume_m3 * 1000, 2)  # m^3 -> litres


def get_tank(fuel_type):
    return db.query_db("SELECT * FROM tanks WHERE fuel_type=?", (fuel_type,), one=True)


def gen_temp_password():
    return secrets.token_urlsafe(6).replace("-", "").replace("_", "")[:8]


# ---------------------------------------------------------------------------
# Page routes
# ---------------------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/healthz")
def healthz():
    """Liveness/readiness probe for load balancers and PaaS platforms.
    Confirms the process can actually reach the database, not just that the
    web process is running."""
    try:
        db.query_db("SELECT 1", one=True)
        return jsonify({"status": "ok"}), 200
    except Exception as exc:  # pragma: no cover - defensive
        return jsonify({"status": "error", "detail": str(exc)}), 503


@app.route("/api/auth/csrf")
def csrf_token():
    return jsonify({"csrf_token": get_csrf_token()})


# ---------------------------------------------------------------------------
# Pump identity + Owner/Manager registration
# ---------------------------------------------------------------------------
@app.route("/api/pump/status")
def pump_status():
    pump = db.query_db("SELECT * FROM pump WHERE id=1", one=True)
    has_owner = bool(db.query_db("SELECT id FROM users WHERE role='admin' LIMIT 1", one=True))
    return jsonify({
        "registered": bool(pump) or has_owner,
        "code_set": bool(pump and pump["pump_code_hash"]),
        "business_name": (pump["business_name"] if pump else get_setting("business_name", "")) or "",
    })


@app.route("/api/auth/register-owner", methods=["POST"])
def register_owner():
    existing_pump = db.query_db("SELECT id FROM pump WHERE id=1", one=True)
    existing_owner = db.query_db("SELECT id FROM users WHERE role='admin' LIMIT 1", one=True)
    if existing_pump or existing_owner:
        return jsonify({"error": "This pump is already registered. Please log in instead."}), 400

    data = request.get_json(force=True, silent=True) or {}
    business_name = (data.get("business_name") or "").strip()
    code = (data.get("pump_code") or "").strip()
    code_confirm = (data.get("pump_code_confirm") or "").strip()
    name = (data.get("name") or "").strip()
    email = (data.get("email") or "").strip().lower()
    password = data.get("password") or ""

    if not business_name:
        return jsonify({"error": "Pump / business name is required."}), 400
    if not PUMP_CODE_RE.match(code):
        return jsonify({"error": "Pump code must be exactly 6 digits."}), 400
    if code != code_confirm:
        return jsonify({"error": "Pump codes do not match."}), 400
    if not name or not email or len(password) < 6:
        return jsonify({"error": "Owner name, email, and a password of 6+ characters are required."}), 400
    if db.query_db("SELECT id FROM users WHERE lower(email)=?", (email,), one=True):
        return jsonify({"error": "A user with this email already exists."}), 400

    avatar = "".join([p[0].upper() for p in name.split()[:2]]) or "OW"
    personal_id = db.next_personal_id("admin")
    uid = db.execute_db(
        "INSERT INTO users (name,email,password_hash,role,avatar,personal_id) VALUES (?,?,?,?,?,?)",
        (name, email, generate_password_hash(password), "admin", avatar, personal_id),
    )
    db.execute_db(
        "INSERT INTO pump (id,business_name,pump_code_hash,owner_user_id) VALUES (1,?,?,?)",
        (business_name, generate_password_hash(code), uid),
    )
    db.execute_db("INSERT OR REPLACE INTO settings (key,value) VALUES ('business_name',?)", (business_name,))

    session.clear()
    session["user_id"] = uid
    session["role"] = "admin"
    session.permanent = True
    user = db.query_db("SELECT id,name,email,role,avatar,personal_id FROM users WHERE id=?", (uid,), one=True)
    audit("register_owner", "pump", uid, f"Pump '{business_name}' registered", actor_override=user)
    return jsonify({**serialize_user(user), "csrf_token": get_csrf_token()}), 201


@app.route("/api/auth/register-manager", methods=["POST"])
def register_manager():
    ip_key = "pumpcode:" + _client_ip()
    locked = is_login_locked(ip_key)
    if locked:
        return jsonify({"error": f"Too many incorrect pump codes. Try again after {locked.strftime('%H:%M')}."}), 429

    pump = db.query_db("SELECT * FROM pump WHERE id=1", one=True)
    if not pump or not pump["pump_code_hash"]:
        return jsonify({"error": "The pump has not completed setup yet. Ask the Owner to set the pump code first."}), 400

    data = request.get_json(force=True, silent=True) or {}
    code = (data.get("pump_code") or "").strip()
    name = (data.get("name") or "").strip()
    email = (data.get("email") or "").strip().lower()
    password = data.get("password") or ""

    if not check_password_hash(pump["pump_code_hash"], code):
        register_failed_login(ip_key)
        audit("register_manager", "pump", None, "Invalid pump code", success=False, actor_override=None)
        return jsonify({"error": "Invalid pump code."}), 401
    clear_failed_logins(ip_key)

    if not name or not email or len(password) < 6:
        return jsonify({"error": "Name, email, and a password of 6+ characters are required."}), 400
    if db.query_db("SELECT id FROM users WHERE lower(email)=?", (email,), one=True):
        return jsonify({"error": "A user with this email already exists."}), 400

    avatar = "".join([p[0].upper() for p in name.split()[:2]]) or "MG"
    personal_id = db.next_personal_id("manager")
    uid = db.execute_db(
        "INSERT INTO users (name,email,password_hash,role,avatar,personal_id) VALUES (?,?,?,?,?,?)",
        (name, email, generate_password_hash(password), "manager", avatar, personal_id),
    )
    session.clear()
    session["user_id"] = uid
    session["role"] = "manager"
    session.permanent = True
    user = db.query_db("SELECT id,name,email,role,avatar,personal_id FROM users WHERE id=?", (uid,), one=True)
    audit("register_manager", "user", uid, f"Manager '{name}' registered via pump code", actor_override=user)
    return jsonify({**serialize_user(user), "csrf_token": get_csrf_token()}), 201


@app.route("/api/pump", methods=["GET"])
@roles_required("admin")
def get_pump():
    pump = db.query_db("SELECT id,business_name,owner_user_id,created_at FROM pump WHERE id=1", one=True)
    if not pump:
        return jsonify({"error": "Pump not found"}), 404
    full = db.query_db("SELECT pump_code_hash FROM pump WHERE id=1", one=True)
    pump["code_set"] = bool(full["pump_code_hash"])
    return jsonify(pump)


@app.route("/api/pump/code", methods=["PUT"])
@roles_required("admin")
def set_pump_code():
    data = request.get_json(force=True, silent=True) or {}
    new_code = (data.get("new_code") or "").strip()
    confirm_code = (data.get("confirm_code") or "").strip()
    current_password = data.get("current_password") or ""

    pump = db.query_db("SELECT * FROM pump WHERE id=1", one=True)
    if not pump:
        return jsonify({"error": "Pump not found"}), 404

    if pump["pump_code_hash"]:
        user = db.query_db("SELECT * FROM users WHERE id=?", (session["user_id"],), one=True)
        if not check_password_hash(user["password_hash"], current_password):
            audit("pump_code_reset", "pump", 1, "Failed: wrong password", success=False)
            return jsonify({"error": "Please confirm your current password to reset the pump code."}), 400

    if not PUMP_CODE_RE.match(new_code):
        return jsonify({"error": "Pump code must be exactly 6 digits."}), 400
    if new_code != confirm_code:
        return jsonify({"error": "Pump codes do not match."}), 400

    was_set = bool(pump["pump_code_hash"])
    db.execute_db("UPDATE pump SET pump_code_hash=? WHERE id=1", (generate_password_hash(new_code),))
    audit("pump_code_reset" if was_set else "pump_code_set", "pump", 1, "Pump code updated by Owner")
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Login / session
# ---------------------------------------------------------------------------
@app.route("/api/auth/login", methods=["POST"])
def login():
    data = request.get_json(force=True, silent=True) or {}
    identifier = (data.get("identifier") or data.get("email") or data.get("personal_id") or "").strip()
    password = data.get("password") or ""
    id_key = identifier.lower()

    locked = is_login_locked(id_key)
    if locked:
        audit("login", "user", None, f"Locked out ({identifier})", success=False)
        return jsonify({"error": f"Too many failed attempts. Try again after {locked.strftime('%H:%M')}."}), 429

    user = db.query_db(
        "SELECT * FROM users WHERE lower(email)=? OR upper(personal_id)=?",
        (id_key, identifier.upper()), one=True,
    )
    if not user or not check_password_hash(user["password_hash"], password) or not user["active"]:
        register_failed_login(id_key)
        db.log_action(user["id"] if user else None, user["name"] if user else None,
                       "login", "user", user["id"] if user else None,
                       f"Failed login for '{identifier}'", _client_ip(), success=False)
        return jsonify({"error": "Invalid credentials."}), 401

    clear_failed_logins(id_key)
    session.clear()
    session["user_id"] = user["id"]
    session["role"] = user["role"]
    session.permanent = True
    audit("login", "user", user["id"], f"Login as {user['role']}",
          actor_override={"id": user["id"], "name": user["name"]})
    return jsonify({
        "id": user["id"], "name": user["name"], "email": user["email"],
        "role": user["role"], "role_label": db.ROLE_LABELS.get(user["role"], user["role"]),
        "avatar": user["avatar"], "personal_id": user["personal_id"],
        "csrf_token": get_csrf_token(),
    })


@app.route("/api/auth/logout", methods=["POST"])
def logout():
    u = current_user()
    if u:
        audit("logout", "user", u["id"])
    session.clear()
    return jsonify({"ok": True})


@app.route("/api/auth/me")
def me():
    user = current_user()
    if not user:
        return jsonify({"user": None})
    return jsonify({"user": serialize_user(user), "csrf_token": get_csrf_token()})


@app.route("/api/auth/password", methods=["PUT"])
@login_required
def change_password():
    data = request.get_json(force=True, silent=True) or {}
    current_pw = data.get("current_password") or ""
    new_pw = data.get("new_password") or ""
    user = db.query_db("SELECT * FROM users WHERE id=?", (session["user_id"],), one=True)
    if not check_password_hash(user["password_hash"], current_pw):
        return jsonify({"error": "Current password is incorrect."}), 400
    if len(new_pw) < 6:
        return jsonify({"error": "New password must be at least 6 characters."}), 400
    db.execute_db("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(new_pw), user["id"]))
    audit("password_change", "user", user["id"])
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------
@app.route("/api/dashboard")
@login_required
def dashboard():
    today = datetime.now().date()
    month_start = today.replace(day=1).isoformat()

    total_udhari = db.query_db("SELECT COALESCE(SUM(amount),0) t FROM entries", one=True)["t"]
    total_pending = db.query_db("SELECT COALESCE(SUM(amount-paid_amount),0) t FROM entries", one=True)["t"]
    total_customers = db.query_db("SELECT COUNT(*) c FROM customers", one=True)["c"]
    total_transactions_month = db.query_db(
        "SELECT COUNT(*) c FROM entries WHERE date>=?", (month_start,), one=True
    )["c"]
    received_month = db.query_db(
        "SELECT COALESCE(SUM(amount),0) t FROM payments WHERE date>=?", (month_start,), one=True
    )["t"]

    overdue_cutoff = (today - timedelta(days=30)).isoformat()
    overdue_amount = db.query_db(
        "SELECT COALESCE(SUM(amount-paid_amount),0) t FROM entries WHERE date<=? AND amount>paid_amount",
        (overdue_cutoff,), one=True,
    )["t"]

    chart_start = (today - timedelta(days=29)).isoformat()
    udhari_by_day = {r["d"]: r["t"] for r in db.query_db(
        "SELECT date d, SUM(amount) t FROM entries WHERE date>=? GROUP BY date", (chart_start,)
    )}
    received_by_day = {r["d"]: r["t"] for r in db.query_db(
        "SELECT date d, SUM(amount) t FROM payments WHERE date>=? GROUP BY date", (chart_start,)
    )}
    chart_labels, chart_udhari, chart_received = [], [], []
    for i in range(30):
        d = (today - timedelta(days=29 - i)).isoformat()
        chart_labels.append(d)
        chart_udhari.append(round(udhari_by_day.get(d, 0), 2))
        chart_received.append(round(received_by_day.get(d, 0), 2))

    recent_entries = db.query_db(
        """SELECT e.id, e.date, c.name customer, e.fuel_type, e.qty, e.amount, e.paid_amount,
                  (e.amount-e.paid_amount) pending, e.status
           FROM entries e JOIN customers c ON c.id=e.customer_id
           ORDER BY e.date DESC, e.id DESC LIMIT 6"""
    )

    pending_by_fuel = db.query_db(
        "SELECT fuel_type, COALESCE(SUM(amount-paid_amount),0) t FROM entries GROUP BY fuel_type"
    )
    pending_total = sum(r["t"] for r in pending_by_fuel) or 1
    pending_by_fuel = [{
        "fuel_type": r["fuel_type"], "amount": round(r["t"], 2),
        "pct": round(r["t"] / pending_total * 100, 1),
    } for r in pending_by_fuel]

    top_customers = db.query_db(
        """SELECT c.id, c.name, COALESCE(SUM(e.amount-e.paid_amount),0) pending
           FROM customers c JOIN entries e ON e.customer_id=c.id
           GROUP BY c.id ORDER BY pending DESC LIMIT 5"""
    )

    stock = db.query_db("SELECT * FROM stock")
    low_stock = [s for s in stock if s["quantity"] <= s["low_threshold"]]

    today_iso = today.isoformat()
    attendance_today = db.query_db(
        "SELECT status, COUNT(*) c FROM attendance WHERE date=? GROUP BY status", (today_iso,)
    )
    staff_total = db.query_db("SELECT COUNT(*) c FROM users WHERE role='staff' AND active=1", one=True)["c"]

    return jsonify({
        "stats": {
            "total_udhari": round(total_udhari, 2),
            "total_customers": total_customers,
            "total_transactions_month": total_transactions_month,
            "received_month": round(received_month, 2),
            "total_pending": round(total_pending, 2),
        },
        "overdue_amount": round(overdue_amount, 2),
        "chart": {"labels": chart_labels, "udhari": chart_udhari, "received": chart_received},
        "recent_entries": recent_entries,
        "pending_by_fuel": pending_by_fuel,
        "top_customers": top_customers,
        "stock": stock,
        "low_stock": low_stock,
        "attendance_today": attendance_today,
        "staff_total": staff_total,
    })


# ---------------------------------------------------------------------------
# Customers
# ---------------------------------------------------------------------------
@app.route("/api/customers", methods=["GET"])
@login_required
def list_customers():
    search = request.args.get("search", "").strip()
    sql = """SELECT c.*, COALESCE(SUM(e.amount-e.paid_amount),0) pending,
                     COALESCE(SUM(e.amount),0) total_udhari, COUNT(e.id) txn_count
              FROM customers c LEFT JOIN entries e ON e.customer_id=c.id"""
    args = ()
    if search:
        sql += " WHERE c.name LIKE ? OR c.phone LIKE ? OR c.vehicle_no LIKE ?"
        args = (f"%{search}%", f"%{search}%", f"%{search}%")
    sql += " GROUP BY c.id ORDER BY c.name"
    rows = db.query_db(sql, args)
    for r in rows:
        r["pending"] = round(r["pending"], 2)
        r["total_udhari"] = round(r["total_udhari"], 2)
    return jsonify(rows)


@app.route("/api/customers/<int:cid>", methods=["GET"])
@login_required
def get_customer(cid):
    customer = db.query_db("SELECT * FROM customers WHERE id=?", (cid,), one=True)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404
    entries = db.query_db(
        "SELECT * FROM entries WHERE customer_id=? ORDER BY date DESC, id DESC", (cid,)
    )
    payments = db.query_db(
        "SELECT * FROM payments WHERE customer_id=? ORDER BY date DESC, id DESC", (cid,)
    )
    pending = sum(e["amount"] - e["paid_amount"] for e in entries)
    return jsonify({"customer": customer, "entries": entries, "payments": payments, "pending": round(pending, 2)})


@app.route("/api/customers", methods=["POST"])
@login_required
def create_customer():
    data = request.get_json(force=True, silent=True) or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "Customer name is required."}), 400
    cid = db.execute_db(
        "INSERT INTO customers (name,phone,address,vehicle_no,gst_no) VALUES (?,?,?,?,?)",
        (name, data.get("phone", ""), data.get("address", ""), data.get("vehicle_no", ""), data.get("gst_no", "")),
    )
    audit("create", "customer", cid, name)
    return jsonify(db.query_db("SELECT * FROM customers WHERE id=?", (cid,), one=True)), 201


@app.route("/api/customers/<int:cid>", methods=["PUT"])
@roles_required("admin", "manager")
def update_customer(cid):
    data = request.get_json(force=True, silent=True) or {}
    existing = db.query_db("SELECT * FROM customers WHERE id=?", (cid,), one=True)
    if not existing:
        return jsonify({"error": "Customer not found"}), 404
    db.execute_db(
        "UPDATE customers SET name=?, phone=?, address=?, vehicle_no=?, gst_no=? WHERE id=?",
        (
            data.get("name", existing["name"]), data.get("phone", existing["phone"]),
            data.get("address", existing["address"]), data.get("vehicle_no", existing["vehicle_no"]),
            data.get("gst_no", existing["gst_no"]), cid,
        ),
    )
    audit("update", "customer", cid)
    return jsonify(db.query_db("SELECT * FROM customers WHERE id=?", (cid,), one=True))


@app.route("/api/customers/<int:cid>", methods=["DELETE"])
@roles_required("admin")
def delete_customer(cid):
    txns = db.query_db("SELECT COUNT(*) c FROM entries WHERE customer_id=?", (cid,), one=True)["c"]
    if txns > 0:
        return jsonify({"error": "Cannot delete a customer with existing udhari records."}), 400
    db.execute_db("DELETE FROM customers WHERE id=?", (cid,))
    audit("delete", "customer", cid)
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Entries (Udhari records)
# ---------------------------------------------------------------------------
def _entries_filtered_query(args):
    where, params = ["1=1"], []
    if args.get("search"):
        where.append("(c.name LIKE ? OR e.vehicle_no LIKE ? OR e.driver_name LIKE ?)")
        s = f"%{args['search']}%"
        params += [s, s, s]
    if args.get("fuel") and args["fuel"] != "all":
        where.append("e.fuel_type=?")
        params.append(args["fuel"])
    if args.get("status") and args["status"] != "all":
        where.append("e.status=?")
        params.append(args["status"])
    if args.get("customer_id"):
        where.append("e.customer_id=?")
        params.append(args["customer_id"])
    if args.get("date_from"):
        where.append("e.date>=?")
        params.append(args["date_from"])
    if args.get("date_to"):
        where.append("e.date<=?")
        params.append(args["date_to"])
    return " AND ".join(where), params


@app.route("/api/entries", methods=["GET"])
@login_required
def list_entries():
    where, params = _entries_filtered_query(request.args)
    page = max(int(request.args.get("page", 1)), 1)
    per_page = min(max(int(request.args.get("per_page", 15)), 1), 200)
    offset = (page - 1) * per_page

    total = db.query_db(
        f"SELECT COUNT(*) c FROM entries e JOIN customers c ON c.id=e.customer_id WHERE {where}",
        params, one=True,
    )["c"]
    rows = db.query_db(
        f"""SELECT e.*, c.name customer_name, (e.amount-e.paid_amount) pending
            FROM entries e JOIN customers c ON c.id=e.customer_id
            WHERE {where} ORDER BY e.date DESC, e.id DESC LIMIT ? OFFSET ?""",
        params + [per_page, offset],
    )
    return jsonify({"rows": rows, "total": total, "page": page, "per_page": per_page})


@app.route("/api/entries/export", methods=["GET"])
@login_required
def export_entries():
    where, params = _entries_filtered_query(request.args)
    rows = db.query_db(
        f"""SELECT e.date, c.name customer_name, e.fuel_type, e.qty, e.rate, e.amount,
                   e.paid_amount, (e.amount-e.paid_amount) pending, e.status, e.vehicle_no, e.driver_name, e.remarks
            FROM entries e JOIN customers c ON c.id=e.customer_id
            WHERE {where} ORDER BY e.date DESC, e.id DESC""",
        params,
    )
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Date", "Customer", "Fuel Type", "Qty (Ltr/Kg)", "Rate", "Amount",
                      "Paid", "Pending", "Status", "Vehicle No", "Driver Name", "Remarks"])
    for r in rows:
        writer.writerow([r["date"], r["customer_name"], r["fuel_type"], r["qty"], r["rate"], r["amount"],
                          r["paid_amount"], r["pending"], r["status"], r["vehicle_no"], r["driver_name"], r["remarks"]])
    mem = io.BytesIO(buf.getvalue().encode("utf-8"))
    return send_file(mem, mimetype="text/csv", as_attachment=True,
                      download_name=f"udhari_records_{datetime.now().date().isoformat()}.csv")


@app.route("/api/entries/<int:eid>", methods=["GET"])
@login_required
def get_entry(eid):
    row = db.query_db(
        """SELECT e.*, c.name customer_name FROM entries e JOIN customers c ON c.id=e.customer_id
           WHERE e.id=?""", (eid,), one=True,
    )
    if not row:
        return jsonify({"error": "Entry not found"}), 404
    return jsonify(row)


@app.route("/api/entries", methods=["POST"])
@login_required
def create_entry():
    data = request.get_json(force=True, silent=True) or {}
    customer_id = data.get("customer_id")
    new_customer_name = (data.get("new_customer_name") or "").strip()
    if not customer_id and new_customer_name:
        customer_id = db.execute_db("INSERT INTO customers (name) VALUES (?)", (new_customer_name,))
    if not customer_id:
        return jsonify({"error": "Customer is required."}), 400

    fuel_type = data.get("fuel_type")
    if fuel_type not in FUEL_TYPES:
        return jsonify({"error": "Valid fuel type is required."}), 400
    try:
        qty = float(data.get("qty"))
        rate = float(data.get("rate"))
    except (TypeError, ValueError):
        return jsonify({"error": "Quantity and rate must be numbers."}), 400
    if qty <= 0 or rate <= 0:
        return jsonify({"error": "Quantity and rate must be greater than zero."}), 400

    amount = round(qty * rate, 2)
    date = data.get("date") or datetime.now().date().isoformat()
    eid = db.execute_db(
        """INSERT INTO entries (customer_id,date,fuel_type,qty,rate,amount,paid_amount,
           vehicle_no,driver_name,remarks,status,created_by)
           VALUES (?,?,?,?,?,?,0,?,?,?,?,?)""",
        (customer_id, date, fuel_type, qty, rate, amount,
         data.get("vehicle_no", ""), data.get("driver_name", ""), data.get("remarks", ""),
         "pending", session["user_id"]),
    )
    stock = db.query_db("SELECT * FROM stock WHERE fuel_type=?", (fuel_type,), one=True)
    if stock:
        db.execute_db("UPDATE stock SET quantity=quantity-? WHERE fuel_type=?", (qty, fuel_type))

    audit("create", "entry", eid, f"{qty} {fuel_type} = {amount}")
    return jsonify(db.query_db("SELECT * FROM entries WHERE id=?", (eid,), one=True)), 201


@app.route("/api/entries/<int:eid>", methods=["PUT"])
@roles_required("admin", "manager")
def update_entry(eid):
    existing = db.query_db("SELECT * FROM entries WHERE id=?", (eid,), one=True)
    if not existing:
        return jsonify({"error": "Entry not found"}), 404
    data = request.get_json(force=True, silent=True) or {}

    fuel_type = data.get("fuel_type", existing["fuel_type"])
    qty = float(data.get("qty", existing["qty"]))
    rate = float(data.get("rate", existing["rate"]))
    amount = round(qty * rate, 2)
    paid_amount = min(existing["paid_amount"], amount)
    status = recompute_status(amount, paid_amount)

    if fuel_type == existing["fuel_type"]:
        delta = qty - existing["qty"]
        db.execute_db("UPDATE stock SET quantity=quantity-? WHERE fuel_type=?", (delta, fuel_type))
    else:
        db.execute_db("UPDATE stock SET quantity=quantity+? WHERE fuel_type=?", (existing["qty"], existing["fuel_type"]))
        db.execute_db("UPDATE stock SET quantity=quantity-? WHERE fuel_type=?", (qty, fuel_type))

    db.execute_db(
        """UPDATE entries SET customer_id=?, date=?, fuel_type=?, qty=?, rate=?, amount=?,
           paid_amount=?, vehicle_no=?, driver_name=?, remarks=?, status=? WHERE id=?""",
        (
            data.get("customer_id", existing["customer_id"]), data.get("date", existing["date"]),
            fuel_type, qty, rate, amount, paid_amount,
            data.get("vehicle_no", existing["vehicle_no"]), data.get("driver_name", existing["driver_name"]),
            data.get("remarks", existing["remarks"]), status, eid,
        ),
    )
    audit("update", "entry", eid)
    return jsonify(db.query_db("SELECT * FROM entries WHERE id=?", (eid,), one=True))


@app.route("/api/entries/<int:eid>", methods=["DELETE"])
@roles_required("admin", "manager")
def delete_entry(eid):
    existing = db.query_db("SELECT * FROM entries WHERE id=?", (eid,), one=True)
    if not existing:
        return jsonify({"error": "Entry not found"}), 404
    db.execute_db("UPDATE stock SET quantity=quantity+? WHERE fuel_type=?", (existing["qty"], existing["fuel_type"]))
    db.execute_db("DELETE FROM payments WHERE entry_id=?", (eid,))
    db.execute_db("DELETE FROM entries WHERE id=?", (eid,))
    audit("delete", "entry", eid)
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Payments
# ---------------------------------------------------------------------------
@app.route("/api/payments", methods=["GET"])
@login_required
def list_payments():
    search = request.args.get("search", "").strip()
    sql = """SELECT p.*, c.name customer_name FROM payments p JOIN customers c ON c.id=p.customer_id"""
    args = ()
    if search:
        sql += " WHERE c.name LIKE ?"
        args = (f"%{search}%",)
    sql += " ORDER BY p.date DESC, p.id DESC"
    return jsonify(db.query_db(sql, args))


@app.route("/api/payments", methods=["POST"])
@login_required
def create_payment():
    data = request.get_json(force=True, silent=True) or {}
    customer_id = data.get("customer_id")
    if not customer_id:
        return jsonify({"error": "Customer is required."}), 400
    try:
        amount = float(data.get("amount"))
    except (TypeError, ValueError):
        return jsonify({"error": "Amount must be a number."}), 400
    if amount <= 0:
        return jsonify({"error": "Amount must be greater than zero."}), 400

    date = data.get("date") or datetime.now().date().isoformat()
    mode = data.get("mode", "cash")
    remarks = data.get("remarks", "")
    entry_id = data.get("entry_id")

    if entry_id:
        entry = db.query_db("SELECT * FROM entries WHERE id=? AND customer_id=?", (entry_id, customer_id), one=True)
        if not entry:
            return jsonify({"error": "Entry not found for this customer."}), 404
        pending = entry["amount"] - entry["paid_amount"]
        applied = min(amount, pending)
        new_paid = entry["paid_amount"] + applied
        db.execute_db(
            "UPDATE entries SET paid_amount=?, status=? WHERE id=?",
            (new_paid, recompute_status(entry["amount"], new_paid), entry["id"]),
        )
        pid = db.execute_db(
            "INSERT INTO payments (customer_id,entry_id,amount,date,mode,remarks,created_by) VALUES (?,?,?,?,?,?,?)",
            (customer_id, entry_id, applied, date, mode, remarks, session["user_id"]),
        )
        leftover = amount - applied
        if leftover > 0.001:
            _allocate_to_oldest_pending(customer_id, leftover, date, mode, remarks)
    else:
        pid = db.execute_db(
            "INSERT INTO payments (customer_id,entry_id,amount,date,mode,remarks,created_by) VALUES (?,NULL,?,?,?,?,?)",
            (customer_id, amount, date, mode, remarks, session["user_id"]),
        )
        _allocate_to_oldest_pending(customer_id, amount, date, mode, remarks)

    audit("create", "payment", pid, f"{amount} for customer #{customer_id}")
    return jsonify(db.query_db("SELECT * FROM payments WHERE id=?", (pid,), one=True)), 201


def _allocate_to_oldest_pending(customer_id, amount, date, mode, remarks):
    """Distribute a lump-sum payment across a customer's oldest pending entries first."""
    remaining = amount
    pending_entries = db.query_db(
        "SELECT * FROM entries WHERE customer_id=? AND amount>paid_amount ORDER BY date ASC, id ASC",
        (customer_id,),
    )
    for entry in pending_entries:
        if remaining <= 0.001:
            break
        pending = entry["amount"] - entry["paid_amount"]
        applied = min(remaining, pending)
        new_paid = entry["paid_amount"] + applied
        db.execute_db(
            "UPDATE entries SET paid_amount=?, status=? WHERE id=?",
            (new_paid, recompute_status(entry["amount"], new_paid), entry["id"]),
        )
        remaining -= applied


@app.route("/api/payments/<int:pid>", methods=["DELETE"])
@roles_required("admin", "manager")
def delete_payment(pid):
    payment = db.query_db("SELECT * FROM payments WHERE id=?", (pid,), one=True)
    if not payment:
        return jsonify({"error": "Payment not found"}), 404
    if payment["entry_id"]:
        entry = db.query_db("SELECT * FROM entries WHERE id=?", (payment["entry_id"],), one=True)
        if entry:
            new_paid = max(entry["paid_amount"] - payment["amount"], 0)
            db.execute_db(
                "UPDATE entries SET paid_amount=?, status=? WHERE id=?",
                (new_paid, recompute_status(entry["amount"], new_paid), entry["id"]),
            )
    db.execute_db("DELETE FROM payments WHERE id=?", (pid,))
    audit("delete", "payment", pid)
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Reports (udhari)
# ---------------------------------------------------------------------------
@app.route("/api/reports/summary")
@login_required
def reports_summary():
    where, params = _entries_filtered_query(request.args)
    totals = db.query_db(
        f"""SELECT COALESCE(SUM(e.amount),0) udhari, COALESCE(SUM(e.paid_amount),0) received,
                   COALESCE(SUM(e.amount-e.paid_amount),0) pending, COUNT(*) txns
            FROM entries e JOIN customers c ON c.id=e.customer_id WHERE {where}""",
        params, one=True,
    )
    by_fuel = db.query_db(
        f"""SELECT e.fuel_type, COALESCE(SUM(e.amount),0) udhari, COALESCE(SUM(e.paid_amount),0) received,
                   COALESCE(SUM(e.amount-e.paid_amount),0) pending, COUNT(*) txns
            FROM entries e JOIN customers c ON c.id=e.customer_id WHERE {where} GROUP BY e.fuel_type""",
        params,
    )
    by_customer = db.query_db(
        f"""SELECT c.name customer_name, COALESCE(SUM(e.amount),0) udhari, COALESCE(SUM(e.paid_amount),0) received,
                   COALESCE(SUM(e.amount-e.paid_amount),0) pending, COUNT(*) txns
            FROM entries e JOIN customers c ON c.id=e.customer_id WHERE {where}
            GROUP BY c.id ORDER BY udhari DESC LIMIT 20""",
        params,
    )
    by_date = db.query_db(
        f"""SELECT e.date, COALESCE(SUM(e.amount),0) udhari, COALESCE(SUM(e.paid_amount),0) received
            FROM entries e JOIN customers c ON c.id=e.customer_id WHERE {where}
            GROUP BY e.date ORDER BY e.date ASC""",
        params,
    )
    return jsonify({"totals": totals, "by_fuel": by_fuel, "by_customer": by_customer, "by_date": by_date})


@app.route("/api/reports/export")
@login_required
def reports_export():
    where, params = _entries_filtered_query(request.args)
    rows = db.query_db(
        f"""SELECT e.date, c.name customer_name, e.fuel_type, e.qty, e.rate, e.amount,
                   e.paid_amount, (e.amount-e.paid_amount) pending, e.status
            FROM entries e JOIN customers c ON c.id=e.customer_id
            WHERE {where} ORDER BY e.date ASC""",
        params,
    )
    fmt = request.args.get("format", "csv")
    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Report"
        ws.append(["Date", "Customer", "Fuel Type", "Qty", "Rate", "Amount", "Paid", "Pending", "Status"])
        for r in rows:
            ws.append([r["date"], r["customer_name"], r["fuel_type"], r["qty"], r["rate"],
                       r["amount"], r["paid_amount"], r["pending"], r["status"]])
        mem = io.BytesIO()
        wb.save(mem)
        mem.seek(0)
        return send_file(mem, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                          as_attachment=True, download_name=f"report_{datetime.now().date().isoformat()}.xlsx")
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Date", "Customer", "Fuel Type", "Qty", "Rate", "Amount", "Paid", "Pending", "Status"])
    for r in rows:
        writer.writerow([r["date"], r["customer_name"], r["fuel_type"], r["qty"], r["rate"],
                          r["amount"], r["paid_amount"], r["pending"], r["status"]])
    mem = io.BytesIO(buf.getvalue().encode("utf-8"))
    return send_file(mem, mimetype="text/csv", as_attachment=True,
                      download_name=f"report_{datetime.now().date().isoformat()}.csv")


# ---------------------------------------------------------------------------
# Upload (bulk Excel import — udhari)
# ---------------------------------------------------------------------------
@app.route("/api/upload/sample")
@login_required
def upload_sample():
    wb = Workbook()
    ws = wb.active
    ws.title = "Udhari Upload"
    headers = ["Customer Name", "Date (YYYY-MM-DD)", "Fuel Type (diesel/petrol/cng)",
               "Quantity", "Rate", "Vehicle No", "Driver Name", "Remarks"]
    ws.append(headers)
    ws.append(["Vikram Transport", datetime.now().date().isoformat(), "diesel", 200, 92, "UP42 AB 1234", "Ramesh", "Sample row"])
    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    return send_file(mem, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                      as_attachment=True, download_name="udhari_upload_sample.xlsx")


@app.route("/api/upload", methods=["POST"])
@roles_required("admin", "manager")
def upload_data():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded."}), 400
    file = request.files["file"]
    if not file.filename.lower().endswith((".xlsx", ".xls")):
        return jsonify({"error": "Please upload an .xlsx or .xls file."}), 400

    try:
        wb = load_workbook(file, data_only=True)
    except Exception:
        return jsonify({"error": "Could not read the Excel file. Please use the sample template."}), 400

    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    imported, errors = 0, []

    for idx, row in enumerate(rows, start=2):
        if row is None or all(c is None for c in row):
            continue
        try:
            name = str(row[0]).strip() if row[0] else ""
            date_val = row[1]
            fuel = str(row[2]).strip().lower() if row[2] else ""
            qty = float(row[3]) if row[3] not in (None, "") else 0
            rate = float(row[4]) if row[4] not in (None, "") else 0
            vehicle_no = str(row[5]).strip() if len(row) > 5 and row[5] else ""
            driver_name = str(row[6]).strip() if len(row) > 6 and row[6] else ""
            remarks = str(row[7]).strip() if len(row) > 7 and row[7] else ""

            if not name:
                errors.append(f"Row {idx}: customer name missing.")
                continue
            if fuel not in FUEL_TYPES:
                errors.append(f"Row {idx}: fuel type '{fuel}' invalid (use diesel/petrol/cng).")
                continue
            if qty <= 0 or rate <= 0:
                errors.append(f"Row {idx}: quantity and rate must be greater than zero.")
                continue

            if hasattr(date_val, "date"):
                date_str = date_val.date().isoformat()
            elif isinstance(date_val, str) and date_val.strip():
                date_str = date_val.strip()[:10]
            else:
                date_str = datetime.now().date().isoformat()

            customer = db.query_db("SELECT id FROM customers WHERE lower(name)=?", (name.lower(),), one=True)
            customer_id = customer["id"] if customer else db.execute_db(
                "INSERT INTO customers (name) VALUES (?)", (name,)
            )
            amount = round(qty * rate, 2)
            db.execute_db(
                """INSERT INTO entries (customer_id,date,fuel_type,qty,rate,amount,paid_amount,
                   vehicle_no,driver_name,remarks,status,created_by)
                   VALUES (?,?,?,?,?,?,0,?,?,?,'pending',?)""",
                (customer_id, date_str, fuel, qty, rate, amount, vehicle_no, driver_name, remarks, session["user_id"]),
            )
            db.execute_db("UPDATE stock SET quantity=quantity-? WHERE fuel_type=?", (qty, fuel))
            imported += 1
        except Exception as e:
            errors.append(f"Row {idx}: {e}")

    audit("bulk_upload", "entry", None, f"{imported} row(s) imported, {len(errors)} error(s)")
    return jsonify({"imported": imported, "errors": errors, "total_rows": len(rows)})


# ---------------------------------------------------------------------------
# Users (owner-only account management, all roles)
# ---------------------------------------------------------------------------
@app.route("/api/users", methods=["GET"])
@roles_required("admin")
def list_users():
    rows = db.query_db(
        "SELECT id,name,email,role,avatar,active,personal_id,created_at FROM users ORDER BY id"
    )
    for r in rows:
        r["role_label"] = db.ROLE_LABELS.get(r["role"], r["role"])
    return jsonify(rows)


@app.route("/api/users", methods=["POST"])
@roles_required("admin")
def create_user():
    data = request.get_json(force=True, silent=True) or {}
    name = (data.get("name") or "").strip()
    email = (data.get("email") or "").strip().lower()
    password = data.get("password") or ""
    role = data.get("role", "staff")
    if not name or not email or len(password) < 6:
        return jsonify({"error": "Name, email, and a password of 6+ characters are required."}), 400
    if role not in ROLE_RANK:
        return jsonify({"error": "Invalid role."}), 400
    if db.query_db("SELECT id FROM users WHERE lower(email)=?", (email,), one=True):
        return jsonify({"error": "A user with this email already exists."}), 400
    avatar = "".join([p[0].upper() for p in name.split()[:2]]) or "US"
    personal_id = db.next_personal_id(role)
    uid = db.execute_db(
        "INSERT INTO users (name,email,password_hash,role,avatar,personal_id) VALUES (?,?,?,?,?,?)",
        (name, email, generate_password_hash(password), role, avatar, personal_id),
    )
    audit("create", "user", uid, f"{name} ({role})")
    return jsonify(db.query_db(
        "SELECT id,name,email,role,avatar,active,personal_id FROM users WHERE id=?", (uid,), one=True
    )), 201


@app.route("/api/users/<int:uid>", methods=["PUT"])
@roles_required("admin")
def update_user(uid):
    existing = db.query_db("SELECT * FROM users WHERE id=?", (uid,), one=True)
    if not existing:
        return jsonify({"error": "User not found"}), 404
    data = request.get_json(force=True, silent=True) or {}

    if existing["role"] == "admin" and data.get("role") and data["role"] != "admin":
        admin_count = db.query_db("SELECT COUNT(*) c FROM users WHERE role='admin'", one=True)["c"]
        if admin_count <= 1:
            return jsonify({"error": "Cannot demote the last remaining Owner."}), 400

    name = data.get("name", existing["name"])
    role = data.get("role", existing["role"])
    active = int(data.get("active", existing["active"]))
    if role not in ROLE_RANK:
        return jsonify({"error": "Invalid role."}), 400

    db.execute_db("UPDATE users SET name=?, role=?, active=? WHERE id=?", (name, role, active, uid))
    if data.get("password"):
        if len(data["password"]) < 6:
            return jsonify({"error": "Password must be at least 6 characters."}), 400
        db.execute_db("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(data["password"]), uid))
    audit("update", "user", uid)
    return jsonify(db.query_db(
        "SELECT id,name,email,role,avatar,active,personal_id FROM users WHERE id=?", (uid,), one=True
    ))


@app.route("/api/users/<int:uid>", methods=["DELETE"])
@roles_required("admin")
def delete_user(uid):
    if uid == session["user_id"]:
        return jsonify({"error": "You cannot delete your own account while logged in."}), 400
    existing = db.query_db("SELECT * FROM users WHERE id=?", (uid,), one=True)
    if not existing:
        return jsonify({"error": "User not found"}), 404
    if existing["role"] == "admin":
        admin_count = db.query_db("SELECT COUNT(*) c FROM users WHERE role='admin'", one=True)["c"]
        if admin_count <= 1:
            return jsonify({"error": "Cannot delete the last remaining Owner."}), 400
    db.execute_db("DELETE FROM users WHERE id=?", (uid,))
    audit("delete", "user", uid)
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Staff management (HR profile fields, owner + manager)
# ---------------------------------------------------------------------------
STAFF_FIELDS = ["employee_code", "phone", "designation", "joining_date", "shift", "emergency_contact", "address", "notes"]


@app.route("/api/staff", methods=["GET"])
@roles_required("admin", "manager")
def list_staff():
    rows = db.query_db(
        """SELECT id,name,email,role,avatar,active,personal_id,employee_code,phone,designation,
                  joining_date,shift,emergency_contact,address,notes,created_at
           FROM users WHERE role='staff' ORDER BY name"""
    )
    return jsonify(rows)


@app.route("/api/staff", methods=["POST"])
@roles_required("admin", "manager")
def create_staff():
    data = request.get_json(force=True, silent=True) or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "Staff name is required."}), 400
    personal_id = db.next_personal_id("staff")
    email = (data.get("email") or "").strip().lower() or f"{personal_id.lower()}@staff.local"
    if db.query_db("SELECT id FROM users WHERE lower(email)=?", (email,), one=True):
        return jsonify({"error": "A user with this email already exists."}), 400

    password = data.get("password") or ""
    temp_password = None
    if len(password) < 6:
        temp_password = gen_temp_password()
        password = temp_password

    avatar = "".join([p[0].upper() for p in name.split()[:2]]) or "ST"
    uid = db.execute_db(
        """INSERT INTO users (name,email,password_hash,role,avatar,personal_id,employee_code,phone,
           designation,joining_date,shift,emergency_contact,address,notes)
           VALUES (?,?,?,'staff',?,?,?,?,?,?,?,?,?,?)""",
        (name, email, generate_password_hash(password), avatar, personal_id,
         data.get("employee_code", ""), data.get("phone", ""), data.get("designation", ""),
         data.get("joining_date") or datetime.now().date().isoformat(), data.get("shift", ""),
         data.get("emergency_contact", ""), data.get("address", ""), data.get("notes", "")),
    )
    audit("create", "staff", uid, f"{name} ({personal_id})")
    result = db.query_db(
        "SELECT id,name,email,role,avatar,active,personal_id,employee_code,phone,designation,"
        "joining_date,shift,emergency_contact,address,notes FROM users WHERE id=?", (uid,), one=True
    )
    if temp_password:
        result["temp_password"] = temp_password
    return jsonify(result), 201


@app.route("/api/staff/<int:uid>", methods=["PUT"])
@roles_required("admin", "manager")
def update_staff(uid):
    existing = db.query_db("SELECT * FROM users WHERE id=? AND role='staff'", (uid,), one=True)
    if not existing:
        return jsonify({"error": "Staff member not found"}), 404
    data = request.get_json(force=True, silent=True) or {}

    fields = {
        "name": data.get("name", existing["name"]),
        "employee_code": data.get("employee_code", existing["employee_code"]),
        "phone": data.get("phone", existing["phone"]),
        "designation": data.get("designation", existing["designation"]),
        "joining_date": data.get("joining_date", existing["joining_date"]),
        "shift": data.get("shift", existing["shift"]),
        "emergency_contact": data.get("emergency_contact", existing["emergency_contact"]),
        "address": data.get("address", existing["address"]),
        "notes": data.get("notes", existing["notes"]),
        "active": int(data.get("active", existing["active"])),
    }
    db.execute_db(
        """UPDATE users SET name=?,employee_code=?,phone=?,designation=?,joining_date=?,shift=?,
           emergency_contact=?,address=?,notes=?,active=? WHERE id=?""",
        (*fields.values(), uid),
    )
    result = {"ok": True}
    if data.get("reset_password"):
        temp_password = gen_temp_password()
        db.execute_db("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(temp_password), uid))
        result["temp_password"] = temp_password
        audit("password_reset", "staff", uid)
    audit("update", "staff", uid)
    result.update(db.query_db(
        "SELECT id,name,email,role,avatar,active,personal_id,employee_code,phone,designation,"
        "joining_date,shift,emergency_contact,address,notes FROM users WHERE id=?", (uid,), one=True
    ))
    return jsonify(result)


@app.route("/api/staff/<int:uid>", methods=["DELETE"])
@roles_required("admin")
def delete_staff(uid):
    existing = db.query_db("SELECT * FROM users WHERE id=? AND role='staff'", (uid,), one=True)
    if not existing:
        return jsonify({"error": "Staff member not found"}), 404
    has_attendance = db.query_db("SELECT COUNT(*) c FROM attendance WHERE staff_id=?", (uid,), one=True)["c"]
    if has_attendance:
        return jsonify({"error": "This staff member has attendance history. Deactivate instead of deleting."}), 400
    db.execute_db("DELETE FROM users WHERE id=?", (uid,))
    audit("delete", "staff", uid)
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Attendance
# ---------------------------------------------------------------------------
@app.route("/api/attendance", methods=["GET"])
@roles_required("admin", "manager")
def get_attendance():
    date = request.args.get("date") or datetime.now().date().isoformat()
    rows = db.query_db(
        """SELECT u.id staff_id, u.name, u.personal_id, u.designation, u.employee_code,
                  a.id attendance_id, a.status, a.check_in, a.check_out, a.shift, a.remarks
           FROM users u LEFT JOIN attendance a ON a.staff_id=u.id AND a.date=?
           WHERE u.role='staff' AND u.active=1 ORDER BY u.name""",
        (date,),
    )
    return jsonify({"date": date, "rows": rows})


@app.route("/api/attendance", methods=["POST"])
@roles_required("admin", "manager")
def mark_attendance():
    data = request.get_json(force=True, silent=True) or {}
    staff_id = data.get("staff_id")
    date = data.get("date") or datetime.now().date().isoformat()
    status = data.get("status", "present")
    if not staff_id:
        return jsonify({"error": "staff_id is required."}), 400
    if status not in ATTENDANCE_STATUSES:
        return jsonify({"error": "Invalid attendance status."}), 400
    staff = db.query_db("SELECT id FROM users WHERE id=? AND role='staff'", (staff_id,), one=True)
    if not staff:
        return jsonify({"error": "Staff member not found."}), 404

    db.execute_db(
        """INSERT INTO attendance (staff_id,date,status,check_in,check_out,shift,remarks,created_by,updated_at)
           VALUES (?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP)
           ON CONFLICT(staff_id,date) DO UPDATE SET
             status=excluded.status, check_in=excluded.check_in, check_out=excluded.check_out,
             shift=excluded.shift, remarks=excluded.remarks, updated_at=CURRENT_TIMESTAMP""",
        (staff_id, date, status, data.get("check_in", ""), data.get("check_out", ""),
         data.get("shift", ""), data.get("remarks", ""), session["user_id"]),
    )
    audit("mark_attendance", "attendance", staff_id, f"{date}: {status}")
    return jsonify({"ok": True})


@app.route("/api/attendance/summary")
@roles_required("admin", "manager")
def attendance_summary():
    month = request.args.get("month") or datetime.now().strftime("%Y-%m")
    rows = db.query_db(
        """SELECT u.id staff_id, u.name, u.personal_id,
                  SUM(CASE WHEN a.status='present' THEN 1 ELSE 0 END) present,
                  SUM(CASE WHEN a.status='absent' THEN 1 ELSE 0 END) absent,
                  SUM(CASE WHEN a.status='half_day' THEN 1 ELSE 0 END) half_day,
                  SUM(CASE WHEN a.status='leave' THEN 1 ELSE 0 END) leave_days,
                  SUM(CASE WHEN a.status='off' THEN 1 ELSE 0 END) off_days
           FROM users u LEFT JOIN attendance a ON a.staff_id=u.id AND substr(a.date,1,7)=?
           WHERE u.role='staff' AND u.active=1 GROUP BY u.id ORDER BY u.name""",
        (month,),
    )
    return jsonify({"month": month, "rows": rows})


# ---------------------------------------------------------------------------
# Tanks (read-only geometry info for the dip live-preview)
# ---------------------------------------------------------------------------
@app.route("/api/tanks")
@login_required
def list_tanks():
    return jsonify(db.query_db("SELECT * FROM tanks"))


# ---------------------------------------------------------------------------
# Fuel Receipts (stock inward)
# ---------------------------------------------------------------------------
@app.route("/api/fuel-receipts", methods=["GET"])
@login_required
def list_fuel_receipts():
    date = request.args.get("date")
    fuel = request.args.get("fuel")
    where, params = ["1=1"], []
    if date:
        where.append("date=?"); params.append(date)
    if fuel and fuel != "all":
        where.append("fuel_type=?"); params.append(fuel)
    rows = db.query_db(
        f"SELECT * FROM fuel_receipts WHERE {' AND '.join(where)} ORDER BY date DESC, id DESC LIMIT 200",
        params,
    )
    return jsonify(rows)


@app.route("/api/fuel-receipts", methods=["POST"])
@login_required
def create_fuel_receipt():
    data = request.get_json(force=True, silent=True) or {}
    fuel_type = data.get("fuel_type")
    if fuel_type not in DIP_FUEL_TYPES:
        return jsonify({"error": "Fuel type must be diesel or petrol."}), 400
    try:
        qty_ltr = float(data.get("qty_ltr"))
    except (TypeError, ValueError):
        return jsonify({"error": "Quantity must be a number."}), 400
    if qty_ltr <= 0:
        return jsonify({"error": "Quantity must be greater than zero."}), 400
    date = data.get("date") or datetime.now().date().isoformat()
    shift = data.get("shift", "morning")
    if shift not in SHIFTS:
        return jsonify({"error": "Invalid shift."}), 400

    rid = db.execute_db(
        """INSERT INTO fuel_receipts (date,shift,fuel_type,qty_ltr,invoice_no,supplier,remarks,created_by)
           VALUES (?,?,?,?,?,?,?,?)""",
        (date, shift, fuel_type, qty_ltr, data.get("invoice_no", ""), data.get("supplier", ""),
         data.get("remarks", ""), session["user_id"]),
    )
    db.execute_db("UPDATE stock SET quantity=quantity+? WHERE fuel_type=?", (qty_ltr, fuel_type))
    audit("create", "fuel_receipt", rid, f"{qty_ltr} L {fuel_type}")
    return jsonify(db.query_db("SELECT * FROM fuel_receipts WHERE id=?", (rid,), one=True)), 201


@app.route("/api/fuel-receipts/<int:rid>", methods=["DELETE"])
@roles_required("admin", "manager")
def delete_fuel_receipt(rid):
    r = db.query_db("SELECT * FROM fuel_receipts WHERE id=?", (rid,), one=True)
    if not r:
        return jsonify({"error": "Receipt not found"}), 404
    db.execute_db("UPDATE stock SET quantity=quantity-? WHERE fuel_type=?", (r["qty_ltr"], r["fuel_type"]))
    db.execute_db("DELETE FROM fuel_receipts WHERE id=?", (rid,))
    audit("delete", "fuel_receipt", rid)
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Daily Fuel Testing (density)
# ---------------------------------------------------------------------------
@app.route("/api/density-entries", methods=["GET"])
@login_required
def list_density_entries():
    date = request.args.get("date")
    fuel = request.args.get("fuel")
    where, params = ["1=1"], []
    if date:
        where.append("date=?"); params.append(date)
    if fuel and fuel != "all":
        where.append("fuel_type=?"); params.append(fuel)
    rows = db.query_db(
        f"SELECT * FROM density_entries WHERE {' AND '.join(where)} ORDER BY date DESC, id DESC LIMIT 200",
        params,
    )
    return jsonify(rows)


@app.route("/api/density-entries", methods=["POST"])
@login_required
def create_density_entry():
    data = request.get_json(force=True, silent=True) or {}
    fuel_type = data.get("fuel_type")
    if fuel_type not in DIP_FUEL_TYPES:
        return jsonify({"error": "Fuel type must be diesel or petrol."}), 400
    try:
        density = float(data.get("density"))
    except (TypeError, ValueError):
        return jsonify({"error": "Density is required."}), 400
    result = data.get("result", "pass")
    if result not in ("pass", "fail", "hold"):
        return jsonify({"error": "Result must be pass, fail, or hold."}), 400
    date = data.get("date") or datetime.now().date().isoformat()
    shift = data.get("shift", "morning")
    if shift not in SHIFTS:
        return jsonify({"error": "Invalid shift."}), 400

    did = db.execute_db(
        """INSERT INTO density_entries (date,shift,fuel_type,testing_qty_ltr,density,temperature,
           water_check,appearance,result,sample_no,tested_by,remarks,created_by)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (date, shift, fuel_type, float(data.get("testing_qty_ltr") or 0), density,
         data.get("temperature"), data.get("water_check", "pass"), data.get("appearance", ""),
         result, data.get("sample_no", ""), data.get("tested_by", ""), data.get("remarks", ""),
         session["user_id"]),
    )
    audit("create", "density_entry", did, f"{fuel_type} density {density} — {result}")
    return jsonify(db.query_db("SELECT * FROM density_entries WHERE id=?", (did,), one=True)), 201


@app.route("/api/density-entries/<int:did>", methods=["DELETE"])
@roles_required("admin", "manager")
def delete_density_entry(did):
    if not db.query_db("SELECT id FROM density_entries WHERE id=?", (did,), one=True):
        return jsonify({"error": "Entry not found"}), 404
    db.execute_db("DELETE FROM density_entries WHERE id=?", (did,))
    audit("delete", "density_entry", did)
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Tank Dip (20 KL)
# ---------------------------------------------------------------------------
@app.route("/api/dip-entries", methods=["GET"])
@login_required
def list_dip_entries():
    date = request.args.get("date")
    fuel = request.args.get("fuel")
    where, params = ["1=1"], []
    if date:
        where.append("date=?"); params.append(date)
    if fuel and fuel != "all":
        where.append("fuel_type=?"); params.append(fuel)
    rows = db.query_db(
        f"SELECT * FROM dip_entries WHERE {' AND '.join(where)} ORDER BY date DESC, id DESC LIMIT 200",
        params,
    )
    return jsonify(rows)


@app.route("/api/dip-entries", methods=["POST"])
@login_required
def create_dip_entry():
    data = request.get_json(force=True, silent=True) or {}
    fuel_type = data.get("fuel_type")
    if fuel_type not in DIP_FUEL_TYPES:
        return jsonify({"error": "Fuel type must be diesel or petrol."}), 400
    try:
        dip_mm = float(data.get("dip_mm"))
    except (TypeError, ValueError):
        return jsonify({"error": "Dip (mm) is required."}), 400
    if dip_mm < 0:
        return jsonify({"error": "Dip cannot be negative."}), 400
    date = data.get("date") or datetime.now().date().isoformat()
    shift = data.get("shift", "morning")
    if shift not in SHIFTS:
        return jsonify({"error": "Invalid shift."}), 400

    tank = get_tank(fuel_type)
    if not tank:
        return jsonify({"error": "Tank configuration not found for this fuel."}), 400
    volume = dip_volume_ltr(dip_mm, tank["diameter_m"], tank["length_m"])
    volume = min(volume, tank["capacity_ltr"])

    stock = db.query_db("SELECT quantity FROM stock WHERE fuel_type=?", (fuel_type,), one=True)
    book_stock = stock["quantity"] if stock else None
    variation = (volume - book_stock) if book_stock is not None else None

    did = db.execute_db(
        """INSERT INTO dip_entries (date,shift,fuel_type,dip_mm,dip_volume_ltr,book_stock_ltr,
           variation_ltr,density,temperature,remarks,created_by)
           VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
        (date, shift, fuel_type, dip_mm, volume, book_stock, variation,
         data.get("density"), data.get("temperature"), data.get("remarks", ""), session["user_id"]),
    )
    audit("create", "dip_entry", did, f"{fuel_type} dip {dip_mm}mm = {volume} L")
    return jsonify(db.query_db("SELECT * FROM dip_entries WHERE id=?", (did,), one=True)), 201


@app.route("/api/dip-entries/<int:did>", methods=["DELETE"])
@roles_required("admin", "manager")
def delete_dip_entry(did):
    if not db.query_db("SELECT id FROM dip_entries WHERE id=?", (did,), one=True):
        return jsonify({"error": "Entry not found"}), 404
    db.execute_db("DELETE FROM dip_entries WHERE id=?", (did,))
    audit("delete", "dip_entry", did)
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Daily Sales (Book Closing = Opening + Receipts - Sales - Testing)
# ---------------------------------------------------------------------------
@app.route("/api/daily-sales", methods=["GET"])
@login_required
def list_daily_sales():
    date = request.args.get("date") or datetime.now().date().isoformat()
    shift = request.args.get("shift")
    where, params = ["date=?"], [date]
    if shift and shift != "all":
        where.append("shift=?"); params.append(shift)
    rows = db.query_db(
        f"SELECT * FROM daily_sales WHERE {' AND '.join(where)} ORDER BY fuel_type, shift", params
    )
    return jsonify(rows)


@app.route("/api/daily-sales/prefill")
@login_required
def prefill_daily_sales():
    date = request.args.get("date") or datetime.now().date().isoformat()
    shift = request.args.get("shift", "morning")
    fuel = request.args.get("fuel")
    if fuel not in DIP_FUEL_TYPES:
        return jsonify({"error": "Fuel type must be diesel or petrol."}), 400

    prev = db.query_db(
        "SELECT book_closing FROM daily_sales WHERE fuel_type=? AND shift=? AND date<? ORDER BY date DESC LIMIT 1",
        (fuel, shift, date), one=True,
    )
    if prev:
        opening_stock = prev["book_closing"]
    else:
        stock = db.query_db("SELECT quantity FROM stock WHERE fuel_type=?", (fuel,), one=True)
        opening_stock = stock["quantity"] if stock else 0

    receipt_ltr = db.query_db(
        "SELECT COALESCE(SUM(qty_ltr),0) t FROM fuel_receipts WHERE date=? AND shift=? AND fuel_type=?",
        (date, shift, fuel), one=True,
    )["t"]
    testing_ltr = db.query_db(
        "SELECT COALESCE(SUM(testing_qty_ltr),0) t FROM density_entries WHERE date=? AND shift=? AND fuel_type=?",
        (date, shift, fuel), one=True,
    )["t"]
    latest_dip = db.query_db(
        "SELECT dip_volume_ltr FROM dip_entries WHERE date=? AND shift=? AND fuel_type=? ORDER BY id DESC LIMIT 1",
        (date, shift, fuel), one=True,
    )
    rate = get_setting(f"rate_{fuel}", "0")

    return jsonify({
        "opening_stock": round(opening_stock, 2),
        "receipt_ltr": round(receipt_ltr, 2),
        "testing_ltr": round(testing_ltr, 2),
        "physical_dip_ltr": latest_dip["dip_volume_ltr"] if latest_dip else None,
        "rate": float(rate) if rate else 0,
    })


@app.route("/api/daily-sales", methods=["POST"])
@roles_required("admin", "manager")
def upsert_daily_sales():
    data = request.get_json(force=True, silent=True) or {}
    fuel_type = data.get("fuel_type")
    if fuel_type not in DIP_FUEL_TYPES:
        return jsonify({"error": "Fuel type must be diesel or petrol."}), 400
    date = data.get("date") or datetime.now().date().isoformat()
    shift = data.get("shift", "morning")
    if shift not in SHIFTS:
        return jsonify({"error": "Invalid shift."}), 400
    try:
        opening_stock = float(data.get("opening_stock", 0))
        receipt_ltr = float(data.get("receipt_ltr", 0))
        sales_ltr = float(data.get("sales_ltr", 0))
        rate = float(data.get("rate", 0))
        testing_ltr = float(data.get("testing_ltr", 0))
    except (TypeError, ValueError):
        return jsonify({"error": "All quantities must be numbers."}), 400
    physical_dip_ltr = data.get("physical_dip_ltr")
    physical_dip_ltr = float(physical_dip_ltr) if physical_dip_ltr not in (None, "") else None

    sales_amount = round(sales_ltr * rate, 2)
    book_closing = round(opening_stock + receipt_ltr - sales_ltr - testing_ltr, 2)
    variance = round(physical_dip_ltr - book_closing, 2) if physical_dip_ltr is not None else None

    db.execute_db(
        """INSERT INTO daily_sales (date,shift,fuel_type,opening_stock,receipt_ltr,sales_ltr,rate,
           sales_amount,testing_ltr,book_closing,physical_dip_ltr,variance,remarks,created_by)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
           ON CONFLICT(date,shift,fuel_type) DO UPDATE SET
             opening_stock=excluded.opening_stock, receipt_ltr=excluded.receipt_ltr,
             sales_ltr=excluded.sales_ltr, rate=excluded.rate, sales_amount=excluded.sales_amount,
             testing_ltr=excluded.testing_ltr, book_closing=excluded.book_closing,
             physical_dip_ltr=excluded.physical_dip_ltr, variance=excluded.variance,
             remarks=excluded.remarks, updated_at=CURRENT_TIMESTAMP""",
        (date, shift, fuel_type, opening_stock, receipt_ltr, sales_ltr, rate, sales_amount,
         testing_ltr, book_closing, physical_dip_ltr, variance, data.get("remarks", ""), session["user_id"]),
    )
    db.execute_db("UPDATE stock SET quantity=? WHERE fuel_type=?", (book_closing, fuel_type))
    audit("upsert", "daily_sales", None, f"{date} {shift} {fuel_type}: closing {book_closing} L")
    row = db.query_db(
        "SELECT * FROM daily_sales WHERE date=? AND shift=? AND fuel_type=?", (date, shift, fuel_type), one=True
    )
    return jsonify(row), 201


@app.route("/api/daily-sales/<int:sid>", methods=["DELETE"])
@roles_required("admin", "manager")
def delete_daily_sales(sid):
    if not db.query_db("SELECT id FROM daily_sales WHERE id=?", (sid,), one=True):
        return jsonify({"error": "Entry not found"}), 404
    db.execute_db("DELETE FROM daily_sales WHERE id=?", (sid,))
    audit("delete", "daily_sales", sid)
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Reports: daily operations export + full data export
# ---------------------------------------------------------------------------
@app.route("/api/reports/daily-export")
@login_required
def daily_export():
    date = request.args.get("date") or datetime.now().date().isoformat()
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Daily Sales"
    ws1.append(["Date", "Shift", "Fuel", "Opening Stock (L)", "Receipts (L)", "Sales (L)", "Rate",
                "Sales Amount", "Testing (L)", "Book Closing (L)", "Physical Dip (L)", "Variance (L)", "Remarks"])
    for r in db.query_db("SELECT * FROM daily_sales WHERE date=? ORDER BY shift, fuel_type", (date,)):
        ws1.append([r["date"], r["shift"], r["fuel_type"], r["opening_stock"], r["receipt_ltr"], r["sales_ltr"],
                    r["rate"], r["sales_amount"], r["testing_ltr"], r["book_closing"], r["physical_dip_ltr"],
                    r["variance"], r["remarks"]])

    ws2 = wb.create_sheet("Daily Testing")
    ws2.append(["Date", "Shift", "Fuel", "Testing Qty (L)", "Density", "Temp (°C)", "Water Check",
                "Appearance", "Result", "Sample No", "Tested By", "Remarks"])
    for r in db.query_db("SELECT * FROM density_entries WHERE date=? ORDER BY shift, fuel_type", (date,)):
        ws2.append([r["date"], r["shift"], r["fuel_type"], r["testing_qty_ltr"], r["density"], r["temperature"],
                    r["water_check"], r["appearance"], r["result"], r["sample_no"], r["tested_by"], r["remarks"]])

    ws3 = wb.create_sheet("Tank Dips")
    ws3.append(["Date", "Shift", "Fuel", "Dip (mm)", "Observed (L)", "Book Stock (L)", "Variance (L)",
                "Density", "Temp (°C)", "Remarks"])
    for r in db.query_db("SELECT * FROM dip_entries WHERE date=? ORDER BY shift, fuel_type", (date,)):
        ws3.append([r["date"], r["shift"], r["fuel_type"], r["dip_mm"], r["dip_volume_ltr"], r["book_stock_ltr"],
                    r["variation_ltr"], r["density"], r["temperature"], r["remarks"]])

    ws4 = wb.create_sheet("Fuel Receipts")
    ws4.append(["Date", "Shift", "Fuel", "Qty (L)", "Invoice No", "Supplier", "Remarks"])
    for r in db.query_db("SELECT * FROM fuel_receipts WHERE date=? ORDER BY shift, fuel_type", (date,)):
        ws4.append([r["date"], r["shift"], r["fuel_type"], r["qty_ltr"], r["invoice_no"], r["supplier"], r["remarks"]])

    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    audit("export", "daily_report", None, date)
    return send_file(mem, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                      as_attachment=True, download_name=f"daily_report_{date}.xlsx")


@app.route("/api/reports/full-export")
@roles_required("admin")
def full_export():
    wb = Workbook()

    ws = wb.active
    ws.title = "Customers"
    ws.append(["ID", "Name", "Phone", "Address", "Vehicle No", "GST No", "Created At"])
    for r in db.query_db("SELECT * FROM customers ORDER BY id"):
        ws.append([r["id"], r["name"], r["phone"], r["address"], r["vehicle_no"], r["gst_no"], r["created_at"]])

    ws2 = wb.create_sheet("Udhari Entries")
    ws2.append(["Date", "Customer", "Fuel", "Qty", "Rate", "Amount", "Paid", "Pending", "Status", "Vehicle No", "Driver"])
    for r in db.query_db(
        "SELECT e.*, c.name customer_name FROM entries e JOIN customers c ON c.id=e.customer_id ORDER BY e.date"
    ):
        ws2.append([r["date"], r["customer_name"], r["fuel_type"], r["qty"], r["rate"], r["amount"],
                    r["paid_amount"], r["amount"] - r["paid_amount"], r["status"], r["vehicle_no"], r["driver_name"]])

    ws3 = wb.create_sheet("Payments")
    ws3.append(["Date", "Customer", "Amount", "Mode", "Remarks"])
    for r in db.query_db(
        "SELECT p.*, c.name customer_name FROM payments p JOIN customers c ON c.id=p.customer_id ORDER BY p.date"
    ):
        ws3.append([r["date"], r["customer_name"], r["amount"], r["mode"], r["remarks"]])

    ws4 = wb.create_sheet("Staff")
    ws4.append(["Personal ID", "Name", "Employee Code", "Phone", "Designation", "Shift", "Joining Date", "Active"])
    for r in db.query_db("SELECT * FROM users WHERE role='staff' ORDER BY name"):
        ws4.append([r["personal_id"], r["name"], r["employee_code"], r["phone"], r["designation"],
                    r["shift"], r["joining_date"], "Yes" if r["active"] else "No"])

    ws5 = wb.create_sheet("Attendance")
    ws5.append(["Date", "Staff", "Personal ID", "Status", "Check In", "Check Out", "Shift", "Remarks"])
    for r in db.query_db(
        "SELECT a.*, u.name staff_name, u.personal_id FROM attendance a JOIN users u ON u.id=a.staff_id ORDER BY a.date"
    ):
        ws5.append([r["date"], r["staff_name"], r["personal_id"], r["status"], r["check_in"], r["check_out"],
                    r["shift"], r["remarks"]])

    ws6 = wb.create_sheet("Daily Sales")
    ws6.append(["Date", "Shift", "Fuel", "Opening", "Receipts", "Sales", "Rate", "Amount", "Testing",
                "Book Closing", "Physical Dip", "Variance"])
    for r in db.query_db("SELECT * FROM daily_sales ORDER BY date"):
        ws6.append([r["date"], r["shift"], r["fuel_type"], r["opening_stock"], r["receipt_ltr"], r["sales_ltr"],
                    r["rate"], r["sales_amount"], r["testing_ltr"], r["book_closing"], r["physical_dip_ltr"], r["variance"]])

    ws7 = wb.create_sheet("Daily Testing")
    ws7.append(["Date", "Shift", "Fuel", "Testing Qty", "Density", "Temp", "Water Check", "Result", "Sample No", "Tested By"])
    for r in db.query_db("SELECT * FROM density_entries ORDER BY date"):
        ws7.append([r["date"], r["shift"], r["fuel_type"], r["testing_qty_ltr"], r["density"], r["temperature"],
                    r["water_check"], r["result"], r["sample_no"], r["tested_by"]])

    ws8 = wb.create_sheet("Tank Dips")
    ws8.append(["Date", "Shift", "Fuel", "Dip (mm)", "Observed (L)", "Book Stock", "Variance"])
    for r in db.query_db("SELECT * FROM dip_entries ORDER BY date"):
        ws8.append([r["date"], r["shift"], r["fuel_type"], r["dip_mm"], r["dip_volume_ltr"], r["book_stock_ltr"], r["variation_ltr"]])

    ws9 = wb.create_sheet("Fuel Receipts")
    ws9.append(["Date", "Shift", "Fuel", "Qty (L)", "Invoice No", "Supplier"])
    for r in db.query_db("SELECT * FROM fuel_receipts ORDER BY date"):
        ws9.append([r["date"], r["shift"], r["fuel_type"], r["qty_ltr"], r["invoice_no"], r["supplier"]])

    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    audit("export", "full_export", None, "Owner full data export")
    return send_file(mem, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                      as_attachment=True, download_name=f"full_export_{datetime.now().date().isoformat()}.xlsx")


# ---------------------------------------------------------------------------
# Audit Log (owner only)
# ---------------------------------------------------------------------------
@app.route("/api/audit-log")
@roles_required("admin")
def get_audit_log():
    page = max(int(request.args.get("page", 1)), 1)
    per_page = min(max(int(request.args.get("per_page", 25)), 1), 100)
    offset = (page - 1) * per_page
    action = request.args.get("action")
    where, params = ["1=1"], []
    if action and action != "all":
        where.append("action=?"); params.append(action)
    total = db.query_db(
        f"SELECT COUNT(*) c FROM audit_log WHERE {' AND '.join(where)}", params, one=True
    )["c"]
    rows = db.query_db(
        f"SELECT * FROM audit_log WHERE {' AND '.join(where)} ORDER BY id DESC LIMIT ? OFFSET ?",
        params + [per_page, offset],
    )
    return jsonify({"rows": rows, "total": total, "page": page, "per_page": per_page})


# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------
@app.route("/api/settings", methods=["GET"])
@login_required
def get_settings():
    rows = db.query_db("SELECT key,value FROM settings")
    settings = {r["key"]: r["value"] for r in rows}
    stock = db.query_db("SELECT * FROM stock")
    pump = db.query_db("SELECT pump_code_hash FROM pump WHERE id=1", one=True)
    return jsonify({"settings": settings, "stock": stock, "pump_code_set": bool(pump and pump["pump_code_hash"])})


@app.route("/api/settings", methods=["PUT"])
@roles_required("admin", "manager")
def update_settings():
    data = request.get_json(force=True, silent=True) or {}
    for key in ["business_name", "business_address", "gst_no", "rate_diesel", "rate_petrol", "rate_cng",
                "currency_symbol", "currency_code", "locale"]:
        if key in data:
            db.execute_db("INSERT OR REPLACE INTO settings (key,value) VALUES (?,?)", (key, str(data[key])))
    if data.get("business_name"):
        db.execute_db("UPDATE pump SET business_name=? WHERE id=1", (data["business_name"],))
    if "stock" in data and isinstance(data["stock"], list):
        for s in data["stock"]:
            if s.get("fuel_type") in FUEL_TYPES:
                db.execute_db(
                    "UPDATE stock SET quantity=?, low_threshold=? WHERE fuel_type=?",
                    (float(s.get("quantity", 0)), float(s.get("low_threshold", 0)), s["fuel_type"]),
                )
    audit("update", "settings")
    return get_settings()


# ---------------------------------------------------------------------------
# Entrypoint
# ---------------------------------------------------------------------------
db.ensure_ready()

if __name__ == "__main__":
    app.run(debug=not IS_PRODUCTION, host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
